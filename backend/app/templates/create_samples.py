"""
Create sample DOCX templates with placeholders.

This script generates sample Word documents that can be used as templates.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.config import settings


def create_bank_guarantee_template(output_path: Path):
    """Create Bank Guarantee template."""
    doc = Document()
    
    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("BANK KAFOLATI")
    run.bold = True
    run.font.size = Pt(16)
    
    # Guarantee number
    doc.add_paragraph()
    p = doc.add_paragraph("Kafolat № {{guarantee_number}}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Date
    p = doc.add_paragraph("«{{issue_date}}» yil")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    doc.add_paragraph()
    
    # Main text
    doc.add_paragraph(
        "{{bank_name}} (keyingi o'rinlarda - Bank), MFO: {{bank_mfo}}, "
        "ushbu Kafolat bilan quyidagilarni tasdiqlaydi:"
    )
    
    doc.add_paragraph(
        "1. {{beneficiary_name}} (keyingi o'rinlarda - Benefitsiar), "
        "manzili: {{beneficiary_address}}, bilan "
        "{{principal_name}} (INN: {{principal_inn}}) o'rtasida tuzilgan "
        "{{contract_number}}-son shartnoma bo'yicha (sana: {{contract_date}}) "
        "majburiyatlarni ta'minlash maqsadida,"
    )
    
    doc.add_paragraph(
        "2. Bank ushbu Kafolat bo'yicha Benefitsiarga "
        "{{guarantee_amount}} ({{guarantee_amount_words}}) so'm miqdorida "
        "to'lov majburiyatini oladi."
    )
    
    doc.add_paragraph(
        "3. Ushbu Kafolat {{expiry_date}} sanasigacha amal qiladi."
    )
    
    doc.add_paragraph()
    doc.add_paragraph("4. Kafolat shartlari:")
    
    items = [
        "Bank Benefitsiarning yozma talabi bo'yicha, Prinsipalning shartnoma "
        "majburiyatlarini bajarmasligi yoki lozim darajada bajarmasligi "
        "haqida dalillar taqdim etilganda to'lovni amalga oshiradi;",
        "Kafolat bo'yicha to'lov talabi Kafolat amal qilish muddati davomida "
        "yoki uni tugatilgandan keyin 15 (o'n besh) kun ichida taqdim etilishi kerak;",
        "Bank to'lov talabini olgandan keyin 5 (besh) ish kuni ichida "
        "to'lovni amalga oshiradi."
    ]
    
    for i, item in enumerate(items, 1):
        doc.add_paragraph(f"   {i}) {item}")
    
    doc.add_paragraph()
    
    # Signatures
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    left_cell = table.rows[0].cells[0]
    left_cell.text = "Bank nomidan:\n\n{{authorized_position}}\n\n\n_________________"
    
    right_cell = table.rows[0].cells[1]
    right_cell.text = "{{authorized_person}}\n\nM.O."
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_payment_order_template(output_path: Path):
    """Create Payment Order template."""
    doc = Document()
    
    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("TO'LOV TOPSHIRIQNOMASI")
    run.bold = True
    run.font.size = Pt(14)
    
    doc.add_paragraph()
    
    # Order info
    p = doc.add_paragraph("№ {{order_number}}                  Sana: {{order_date}}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Payer section
    p = doc.add_paragraph()
    run = p.add_run("TO'LOVCHI:")
    run.bold = True
    
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    
    rows_data = [
        ("Nomi:", "{{payer_name}}"),
        ("INN:", "{{payer_inn}}"),
        ("Hisob raqami:", "{{payer_account}}"),
        ("Bank va MFO:", "{{payer_bank}}, MFO: {{payer_mfo}}"),
    ]
    
    for i, (label, value) in enumerate(rows_data):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
    
    doc.add_paragraph()
    
    # Recipient section
    p = doc.add_paragraph()
    run = p.add_run("OLUVCHI:")
    run.bold = True
    
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    
    rows_data = [
        ("Nomi:", "{{recipient_name}}"),
        ("INN:", "{{recipient_inn}}"),
        ("Hisob raqami:", "{{recipient_account}}"),
        ("Bank va MFO:", "{{recipient_bank}}, MFO: {{recipient_mfo}}"),
    ]
    
    for i, (label, value) in enumerate(rows_data):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
    
    doc.add_paragraph()
    
    # Amount
    p = doc.add_paragraph()
    run = p.add_run("SUMMA: ")
    run.bold = True
    p.add_run("{{amount}} so'm ({{amount_words}})")
    
    # Purpose
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("To'lov maqsadi: ")
    run.bold = True
    p.add_run("{{payment_purpose}}")
    
    if "{{payment_code}}":
        doc.add_paragraph(f"To'lov kodi: {{payment_code}}")
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Signatures
    p = doc.add_paragraph("Rahbar: ___________________")
    p = doc.add_paragraph("Bosh hisobchi: ___________________")
    
    doc.add_paragraph()
    p = doc.add_paragraph("M.O.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_credit_agreement_template(output_path: Path):
    """Create Credit Agreement template."""
    doc = Document()
    
    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("KREDIT SHARTNOMASI")
    run.bold = True
    run.font.size = Pt(16)
    
    # Agreement info
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("№ {{agreement_number}}")
    
    doc.add_paragraph()
    
    p = doc.add_paragraph("{{agreement_date}}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    doc.add_paragraph()
    
    # Parties
    doc.add_paragraph(
        "{{bank_name}} (keyingi o'rinlarda «Bank» deb yuritiladi), "
        "MFO: {{bank_mfo}}, bir tomondan,"
    )
    
    doc.add_paragraph("va")
    
    doc.add_paragraph(
        "{{borrower_name}} (keyingi o'rinlarda «Qarz oluvchi» deb yuritiladi), "
        "INN: {{borrower_inn}}, manzili: {{borrower_address}}, "
        "rahbari {{borrower_director}} nomidan faoliyat yurituvchi, "
        "ikkinchi tomondan,"
    )
    
    doc.add_paragraph()
    doc.add_paragraph("quyidagilar to'g'risida ushbu shartnomani tuzdilar:")
    
    doc.add_paragraph()
    
    # Article 1
    p = doc.add_paragraph()
    run = p.add_run("1. SHARTNOMA PREDMETI")
    run.bold = True
    
    doc.add_paragraph(
        "1.1. Bank Qarz oluvchiga {{credit_amount}} ({{credit_amount_words}}) so'm "
        "miqdorida kredit ajratishni majburiyatiga oladi."
    )
    
    doc.add_paragraph(
        "1.2. Kredit maqsadi: {{credit_purpose}}"
    )
    
    doc.add_paragraph(
        f"1.3. Kredit muddati: {{credit_term_months}} oy."
    )
    
    doc.add_paragraph()
    
    # Article 2
    p = doc.add_paragraph()
    run = p.add_run("2. FOIZ STAVKASI VA TO'LOV")
    run.bold = True
    
    doc.add_paragraph(
        "2.1. Kredit bo'yicha yillik foiz stavkasi: {{interest_rate}}%."
    )
    
    doc.add_paragraph(
        "2.2. Kredit va foizlarni qaytarish jadvali: {{repayment_schedule}}"
    )
    
    doc.add_paragraph()
    
    # Article 3
    p = doc.add_paragraph()
    run = p.add_run("3. GAROV")
    run.bold = True
    
    doc.add_paragraph(
        "3.1. Kredit majburiyatlari quyidagi garov bilan ta'minlanadi: "
        "{{collateral_description}}"
    )
    
    doc.add_paragraph()
    
    # Signatures
    p = doc.add_paragraph()
    run = p.add_run("TOMONLARNING REKVIZITLARI:")
    run.bold = True
    
    table = doc.add_table(rows=1, cols=2)
    
    left_cell = table.rows[0].cells[0]
    left_cell.text = (
        "BANK:\n"
        "{{bank_name}}\n"
        "MFO: {{bank_mfo}}\n"
        "H/r: {{bank_account}}\n\n"
        "_________________\n"
        "M.O."
    )
    
    right_cell = table.rows[0].cells[1]
    right_cell.text = (
        "QARZ OLUVCHI:\n"
        "{{borrower_name}}\n"
        "INN: {{borrower_inn}}\n"
        "Manzil: {{borrower_address}}\n\n"
        "_________________\n"
        "{{borrower_director}}\n"
        "M.O."
    )
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_employment_contract_template(output_path: Path):
    """Create Employment Contract template."""
    doc = Document()
    
    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("MEHNAT SHARTNOMASI")
    run.bold = True
    run.font.size = Pt(16)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("№ {{contract_number}}")
    
    doc.add_paragraph()
    
    p = doc.add_paragraph("{{contract_date}}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    doc.add_paragraph()
    
    # Parties
    doc.add_paragraph(
        "{{company_name}} (keyingi o'rinlarda «Ish beruvchi» deb yuritiladi), "
        "rahbari {{company_director}} nomidan bir tomondan,"
    )
    doc.add_paragraph("va")
    doc.add_paragraph(
        "{{employee_name}} (keyingi o'rinlarda «Xodim» deb yuritiladi), "
        "passport: {{employee_passport}}, manzili: {{employee_address}}, "
        "ikkinchi tomondan,"
    )
    
    doc.add_paragraph()
    doc.add_paragraph("quyidagilar to'g'risida kelishib oldilar:")
    doc.add_paragraph()
    
    # Article 1
    p = doc.add_paragraph()
    run = p.add_run("1. SHARTNOMA PREDMETI")
    run.bold = True
    
    doc.add_paragraph(
        "1.1. Ish beruvchi Xodimni {{position}} lavozimiga qabul qiladi."
    )
    doc.add_paragraph(
        "1.2. Xodim {{department}} bo'limida ishlaydi."
    )
    doc.add_paragraph(
        "1.3. Ish boshlash sanasi: {{start_date}}"
    )
    doc.add_paragraph(
        "1.4. Shartnoma muddati: {{contract_term}}"
    )
    
    if "{{probation_period}}":
        doc.add_paragraph(
            "1.5. Sinov muddati: {{probation_period}} kun."
        )
    
    doc.add_paragraph()
    
    # Article 2
    p = doc.add_paragraph()
    run = p.add_run("2. MEHNAT HAQI")
    run.bold = True
    
    doc.add_paragraph(
        "2.1. Xodimning oylik ish haqi: {{salary}} so'm."
    )
    
    doc.add_paragraph()
    
    # Signatures
    table = doc.add_table(rows=1, cols=2)
    
    left_cell = table.rows[0].cells[0]
    left_cell.text = (
        "ISH BERUVCHI:\n"
        "{{company_name}}\n\n"
        "_________________\n"
        "{{company_director}}\n"
        "M.O."
    )
    
    right_cell = table.rows[0].cells[1]
    right_cell.text = (
        "XODIM:\n"
        "{{employee_name}}\n\n"
        "_________________\n"
        f"Tel: {{employee_phone}}"
    )
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_service_contract_template(output_path: Path):
    """Create Service Contract template."""
    doc = Document()
    
    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("XIZMAT KO'RSATISH SHARTNOMASI")
    run.bold = True
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("№ {{contract_number}}")
    
    doc.add_paragraph()
    p = doc.add_paragraph("{{contract_date}}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    doc.add_paragraph()
    
    # Parties
    doc.add_paragraph(
        "{{client_name}} (INN: {{client_inn}}), manzili: {{client_address}}, "
        "rahbari {{client_director}} nomidan (keyingi o'rinlarda «Buyurtmachi»),"
    )
    doc.add_paragraph("va")
    doc.add_paragraph(
        "{{provider_name}} (INN: {{provider_inn}}), manzili: {{provider_address}}, "
        "rahbari {{provider_director}} nomidan (keyingi o'rinlarda «Ijrochi»),"
    )
    
    doc.add_paragraph()
    doc.add_paragraph("quyidagi shartnomani tuzdilar:")
    doc.add_paragraph()
    
    # Article 1
    p = doc.add_paragraph()
    run = p.add_run("1. SHARTNOMA PREDMETI")
    run.bold = True
    
    doc.add_paragraph(
        "1.1. Ijrochi Buyurtmachiga quyidagi xizmatlarni ko'rsatishni majburiyatiga oladi:"
    )
    doc.add_paragraph("{{service_description}}")
    
    doc.add_paragraph()
    
    # Article 2
    p = doc.add_paragraph()
    run = p.add_run("2. SHARTNOMA NARXI VA TO'LOV TARTIBI")
    run.bold = True
    
    doc.add_paragraph(
        "2.1. Shartnoma bo'yicha xizmatlar narxi: {{contract_amount}} so'm."
    )
    doc.add_paragraph("2.2. To'lov shartlari: {{payment_terms}}")
    
    doc.add_paragraph()
    
    # Article 3
    p = doc.add_paragraph()
    run = p.add_run("3. MUDDATLAR")
    run.bold = True
    
    doc.add_paragraph("3.1. Shartnoma boshlanish sanasi: {{start_date}}")
    doc.add_paragraph("3.2. Shartnoma tugash sanasi: {{end_date}}")
    
    doc.add_paragraph()
    
    # Signatures
    table = doc.add_table(rows=1, cols=2)
    
    left_cell = table.rows[0].cells[0]
    left_cell.text = (
        "BUYURTMACHI:\n"
        "{{client_name}}\n"
        "INN: {{client_inn}}\n\n"
        "_________________\n"
        "{{client_director}}\n"
        "M.O."
    )
    
    right_cell = table.rows[0].cells[1]
    right_cell.text = (
        "IJROCHI:\n"
        "{{provider_name}}\n"
        "INN: {{provider_inn}}\n\n"
        "_________________\n"
        "{{provider_director}}\n"
        "M.O."
    )
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Created: {output_path}")


def create_all_templates():
    """Create all sample templates."""
    base_dir = Path(settings.upload_dir) / "templates"
    
    # Bank templates
    create_bank_guarantee_template(base_dir / "bank" / "bank_guarantee.docx")
    create_payment_order_template(base_dir / "bank" / "payment_order.docx")
    create_credit_agreement_template(base_dir / "bank" / "credit_agreement.docx")
    
    # HR templates
    create_employment_contract_template(base_dir / "hr" / "employment_contract.docx")
    
    # Contract templates
    create_service_contract_template(base_dir / "contracts" / "service_contract.docx")
    
    print("\nAll templates created successfully!")


if __name__ == "__main__":
    create_all_templates()
