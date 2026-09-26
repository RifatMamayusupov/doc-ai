"""Document processing service for analyzing and extracting data."""

import asyncio
import json
from pathlib import Path
from typing import Any

from app.config import settings


class DocumentProcessor:
    """Process various document types for data extraction."""

    SUPPORTED_FORMATS = {
        '.pdf': 'pdf',
        '.doc': 'docx',
        '.docx': 'docx',
        '.xls': 'excel',
        '.xlsx': 'excel',
        '.csv': 'csv',
        '.txt': 'text',
        '.json': 'json',
        '.jpg': 'image',
        '.jpeg': 'image',
        '.png': 'image',
    }

    def __init__(self, user_id: str) -> None:
        """Initialize processor for a user."""
        self.user_id = user_id
        self.upload_dir = settings.upload_dir / "users" / user_id

    async def analyze_document(
        self,
        file_path: Path,
        on_progress: callable = None,
    ) -> dict[str, Any]:
        """
        Analyze a document and extract structured data.
        
        Args:
            file_path: Path to the document
            on_progress: Callback for progress updates (progress: int, message: str)
            
        Returns:
            Dictionary with extracted data and metadata
        """
        ext = file_path.suffix.lower()
        doc_type = self.SUPPORTED_FORMATS.get(ext, 'unknown')

        if on_progress:
            await on_progress(10, f"Loading {doc_type} document...")

        result = {
            'file_name': file_path.name,
            'file_type': doc_type,
            'file_size': file_path.stat().st_size,
            'extracted_data': {},
            'metadata': {},
        }

        try:
            if doc_type == 'pdf':
                result = await self._process_pdf(file_path, on_progress)
            elif doc_type == 'docx':
                result = await self._process_docx(file_path, on_progress)
            elif doc_type == 'excel':
                result = await self._process_excel(file_path, on_progress)
            elif doc_type == 'csv':
                result = await self._process_csv(file_path, on_progress)
            elif doc_type == 'image':
                result = await self._process_image(file_path, on_progress)
            elif doc_type == 'text':
                result = await self._process_text(file_path, on_progress)
            elif doc_type == 'json':
                result = await self._process_json(file_path, on_progress)
        except Exception as e:
            result['error'] = str(e)

        if on_progress:
            await on_progress(100, "Processing complete")

        return result

    async def _process_pdf(self, file_path: Path, on_progress: callable) -> dict:
        """Process PDF documents."""
        try:
            import pdfplumber
        except ImportError:
            return {'error': 'pdfplumber not installed'}

        if on_progress:
            await on_progress(20, "Extracting PDF text...")

        result = {
            'file_name': file_path.name,
            'file_type': 'pdf',
            'pages': [],
            'tables': [],
            'text': '',
        }

        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                if on_progress:
                    progress = 20 + int((i / total_pages) * 60)
                    await on_progress(progress, f"Processing page {i+1}/{total_pages}")

                page_text = page.extract_text() or ''
                page_tables = page.extract_tables()

                result['pages'].append({
                    'page_number': i + 1,
                    'text': page_text,
                    'table_count': len(page_tables),
                })
                result['text'] += page_text + '\n'

                for table in page_tables:
                    result['tables'].append({
                        'page': i + 1,
                        'data': table,
                    })

        return result

    async def _process_docx(self, file_path: Path, on_progress: callable) -> dict:
        """Process Word documents."""
        try:
            from docx import Document
        except ImportError:
            return {'error': 'python-docx not installed'}

        if on_progress:
            await on_progress(30, "Reading Word document...")

        doc = Document(file_path)

        result = {
            'file_name': file_path.name,
            'file_type': 'docx',
            'paragraphs': [],
            'tables': [],
            'text': '',
        }

        for para in doc.paragraphs:
            result['paragraphs'].append(para.text)
            result['text'] += para.text + '\n'

        if on_progress:
            await on_progress(60, "Extracting tables...")

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_data.append(row_data)
            result['tables'].append(table_data)

        return result

    async def _process_excel(self, file_path: Path, on_progress: callable) -> dict:
        """Process Excel files."""
        try:
            import openpyxl
        except ImportError:
            return {'error': 'openpyxl not installed'}

        if on_progress:
            await on_progress(20, "Loading Excel workbook...")

        wb = openpyxl.load_workbook(file_path, data_only=True)

        result = {
            'file_name': file_path.name,
            'file_type': 'excel',
            'sheets': {},
            'sheet_names': wb.sheetnames,
        }

        total_sheets = len(wb.sheetnames)
        for i, sheet_name in enumerate(wb.sheetnames):
            if on_progress:
                progress = 20 + int((i / total_sheets) * 60)
                await on_progress(progress, f"Processing sheet: {sheet_name}")

            sheet = wb[sheet_name]
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append(list(row))

            result['sheets'][sheet_name] = {
                'data': data,
                'rows': len(data),
                'cols': len(data[0]) if data else 0,
            }

        return result

    async def _process_csv(self, file_path: Path, on_progress: callable) -> dict:
        """Process CSV files."""
        import csv

        if on_progress:
            await on_progress(30, "Reading CSV data...")

        result = {
            'file_name': file_path.name,
            'file_type': 'csv',
            'headers': [],
            'data': [],
            'row_count': 0,
        }

        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i == 0:
                    result['headers'] = row
                else:
                    result['data'].append(row)
                result['row_count'] = i + 1

        return result

    async def _process_image(self, file_path: Path, on_progress: callable) -> dict:
        """Process images with OCR."""
        result = {
            'file_name': file_path.name,
            'file_type': 'image',
            'text': '',
            'ocr_available': False,
        }

        if on_progress:
            await on_progress(30, "Analyzing image...")

        try:
            from PIL import Image
            img = Image.open(file_path)
            result['dimensions'] = {'width': img.width, 'height': img.height}
            result['mode'] = img.mode

            # Try OCR
            try:
                import pytesseract
                if on_progress:
                    await on_progress(50, "Running OCR...")
                result['text'] = pytesseract.image_to_string(img)
                result['ocr_available'] = True
            except ImportError:
                result['ocr_available'] = False
        except Exception as e:
            result['error'] = str(e)

        return result

    async def _process_text(self, file_path: Path, on_progress: callable) -> dict:
        """Process text files."""
        if on_progress:
            await on_progress(50, "Reading text file...")

        content = file_path.read_text(encoding='utf-8')

        return {
            'file_name': file_path.name,
            'file_type': 'text',
            'text': content,
            'line_count': len(content.splitlines()),
            'char_count': len(content),
        }

    async def _process_json(self, file_path: Path, on_progress: callable) -> dict:
        """Process JSON files."""
        if on_progress:
            await on_progress(50, "Parsing JSON...")

        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        return {
            'file_name': file_path.name,
            'file_type': 'json',
            'data': data,
            'keys': list(data.keys()) if isinstance(data, dict) else None,
        }


class TemplateProcessor:
    """Match and fill document templates."""

    def __init__(self) -> None:
        """Initialize template processor."""
        self.templates_dir = settings.upload_dir / "templates"

    async def match_templates(
        self,
        extracted_data: dict,
        organization: str | None = None,
    ) -> list[dict]:
        """
        Find templates that match the extracted data.
        
        Args:
            extracted_data: Data extracted from user's document
            organization: Optional filter by organization
            
        Returns:
            List of matching templates with scores
        """
        from sqlalchemy import select
        from app.database import async_session_maker
        from app.models.template import Template

        async with async_session_maker() as db:
            query = select(Template).where(Template.is_active == True)
            if organization:
                query = query.where(Template.organization == organization)

            result = await db.execute(query)
            templates = result.scalars().all()

        matches = []
        for template in templates:
            score = self._calculate_match_score(extracted_data, template.schema)
            if score > 0.3:  # Minimum match threshold
                matches.append({
                    'id': template.id,
                    'name': template.name,
                    'organization': template.organization,
                    'category': template.category,
                    'preview_image': template.preview_image,
                    'score': score,
                })

        # Sort by score descending
        matches.sort(key=lambda x: x['score'], reverse=True)
        return matches[:5]  # Top 5 matches

    def _calculate_match_score(
        self,
        extracted_data: dict,
        template_schema: dict | None,
    ) -> float:
        """Calculate how well extracted data matches template schema."""
        if not template_schema or 'fields' not in template_schema:
            return 0.5  # Default score for templates without schema

        required_fields = [
            f['name'] for f in template_schema['fields'] if f.get('required')
        ]

        # Check text content for field keywords
        text = extracted_data.get('text', '') or ''
        text_lower = text.lower()

        matched = 0
        for field in required_fields:
            if field.lower() in text_lower:
                matched += 1

        if len(required_fields) == 0:
            return 0.5

        return matched / len(required_fields)

    async def fill_template(
        self,
        template_id: str,
        data: dict[str, Any],
        on_progress: callable = None,
    ) -> Path:
        """
        Fill a template with provided data.
        
        Args:
            template_id: Template ID
            data: Data to fill into template
            on_progress: Progress callback
            
        Returns:
            Path to filled document
        """
        from sqlalchemy import select
        from app.database import async_session_maker
        from app.models.template import Template

        if on_progress:
            await on_progress(10, "Loading template...")

        async with async_session_maker() as db:
            result = await db.execute(
                select(Template).where(Template.id == template_id)
            )
            template = result.scalar_one_or_none()

        if not template:
            raise ValueError(f"Template {template_id} not found")

        template_path = Path(template.file_path)
        if not template_path.exists():
            raise FileNotFoundError(f"Template file not found: {template_path}")

        if on_progress:
            await on_progress(30, "Processing template...")

        # Determine output path
        output_dir = settings.upload_dir / "output"
        output_dir.mkdir(exist_ok=True)
        output_path = output_dir / f"filled_{template_path.name}"

        # Fill based on template type
        ext = template_path.suffix.lower()
        if ext == '.docx':
            await self._fill_docx_template(template_path, output_path, data, on_progress)
        elif ext in ('.xlsx', '.xls'):
            await self._fill_excel_template(template_path, output_path, data, on_progress)
        else:
            raise ValueError(f"Unsupported template format: {ext}")

        return output_path

    async def _fill_docx_template(
        self,
        template_path: Path,
        output_path: Path,
        data: dict,
        on_progress: callable,
    ) -> None:
        """Fill a Word template with placeholders."""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx required for Word templates")

        if on_progress:
            await on_progress(50, "Filling Word template...")

        doc = Document(template_path)

        # Replace placeholders in paragraphs
        for para in doc.paragraphs:
            for key, value in data.items():
                placeholder = f"{{{{{key}}}}}"  # {{key}}
                if placeholder in para.text:
                    para.text = para.text.replace(placeholder, str(value))

        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for key, value in data.items():
                        placeholder = f"{{{{{key}}}}}"
                        if placeholder in cell.text:
                            cell.text = cell.text.replace(placeholder, str(value))

        if on_progress:
            await on_progress(80, "Saving document...")

        doc.save(output_path)

    async def _fill_excel_template(
        self,
        template_path: Path,
        output_path: Path,
        data: dict,
        on_progress: callable,
    ) -> None:
        """Fill an Excel template."""
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl required for Excel templates")

        if on_progress:
            await on_progress(50, "Filling Excel template...")

        wb = openpyxl.load_workbook(template_path)

        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        for key, value in data.items():
                            placeholder = f"{{{{{key}}}}}"
                            if placeholder in cell.value:
                                cell.value = cell.value.replace(placeholder, str(value))

        if on_progress:
            await on_progress(80, "Saving workbook...")

        wb.save(output_path)
