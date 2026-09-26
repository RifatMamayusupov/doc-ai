"""
Template Service - Core template processing and document generation.

Handles:
- Template matching based on extracted data
- Template filling with field values
- Document generation from templates
"""

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Pt
from openpyxl import load_workbook

from app.config import settings


class TemplateField:
    """Represents a fillable field in a template."""
    
    def __init__(
        self,
        name: str,
        field_type: str = "string",
        required: bool = False,
        default: Any = None,
        description: str = "",
        options: list[str] | None = None,
    ):
        self.name = name
        self.field_type = field_type
        self.required = required
        self.default = default
        self.description = description
        self.options = options

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "type": self.field_type,
            "required": self.required,
            "default": self.default,
            "description": self.description,
            "options": self.options,
        }


class TemplateService:
    """Service for template processing and document generation."""
    
    # Field pattern for placeholders: {{field_name}} or {field_name}
    FIELD_PATTERN = re.compile(r'\{\{?\s*(\w+)\s*\}?\}')
    
    def __init__(self):
        self.templates_dir = Path(settings.upload_dir) / "templates"
        self.output_dir = Path(settings.upload_dir) / "generated"
        self.templates_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)
    
    def extract_fields_from_docx(self, file_path: Path) -> list[TemplateField]:
        """
        Extract field placeholders from a DOCX template.
        
        Looks for patterns like {{company_name}} or {date}.
        """
        fields = {}
        
        try:
            doc = DocxDocument(file_path)
            
            # Check paragraphs
            for para in doc.paragraphs:
                for match in self.FIELD_PATTERN.finditer(para.text):
                    field_name = match.group(1)
                    if field_name not in fields:
                        fields[field_name] = self._infer_field_type(field_name)
            
            # Check tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for match in self.FIELD_PATTERN.finditer(cell.text):
                            field_name = match.group(1)
                            if field_name not in fields:
                                fields[field_name] = self._infer_field_type(field_name)
            
        except Exception as e:
            print(f"Error extracting fields: {e}")
        
        return [
            TemplateField(
                name=name,
                field_type=type_info["type"],
                required=type_info.get("required", False),
                description=type_info.get("description", ""),
            )
            for name, type_info in fields.items()
        ]
    
    def _infer_field_type(self, field_name: str) -> dict:
        """Infer field type from field name."""
        name_lower = field_name.lower()
        
        # Date fields
        if any(x in name_lower for x in ["date", "sana", "kun", "sanasi"]):
            return {"type": "date", "required": True, "description": "Sana"}
        
        # Number/Amount fields
        if any(x in name_lower for x in ["amount", "summa", "sum", "number", "miqdor", "qiymat"]):
            return {"type": "number", "required": True, "description": "Summa"}
        
        # Phone fields
        if any(x in name_lower for x in ["phone", "telefon", "tel"]):
            return {"type": "phone", "required": False, "description": "Telefon raqami"}
        
        # Email fields
        if "email" in name_lower:
            return {"type": "email", "required": False, "description": "Email"}
        
        # Required name fields
        if any(x in name_lower for x in ["company", "korxona", "firma", "name", "ism", "nomi"]):
            return {"type": "string", "required": True, "description": "Nom"}
        
        # Address fields
        if any(x in name_lower for x in ["address", "manzil"]):
            return {"type": "text", "required": False, "description": "Manzil"}
        
        return {"type": "string", "required": False, "description": ""}
    
    def fill_docx_template(
        self,
        template_path: Path,
        field_values: dict[str, Any],
        output_path: Path | None = None,
    ) -> Path:
        """
        Fill a DOCX template with provided field values.
        
        Args:
            template_path: Path to the template file
            field_values: Dictionary of field names to values
            output_path: Optional output path (auto-generated if not provided)
            
        Returns:
            Path to the generated document
        """
        doc = DocxDocument(template_path)
        
        # Process paragraphs
        for para in doc.paragraphs:
            self._replace_fields_in_runs(para, field_values)
        
        # Process tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._replace_fields_in_runs(para, field_values)
        
        # Generate output path if not provided
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_name = f"{template_path.stem}_{timestamp}.docx"
            output_path = self.output_dir / output_name
        
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        
        return output_path
    
    def _replace_fields_in_runs(self, paragraph, field_values: dict):
        """Replace field placeholders in paragraph runs."""
        full_text = paragraph.text
        
        # Find all field matches
        for match in self.FIELD_PATTERN.finditer(full_text):
            field_name = match.group(1)
            placeholder = match.group(0)
            
            if field_name in field_values:
                value = self._format_value(field_values[field_name], field_name)
                full_text = full_text.replace(placeholder, str(value))
        
        # Clear and rebuild paragraph if changes were made
        if full_text != paragraph.text:
            # Get original formatting
            if paragraph.runs:
                first_run = paragraph.runs[0]
                font_name = first_run.font.name
                font_size = first_run.font.size
                bold = first_run.font.bold
                italic = first_run.font.italic
            else:
                font_name = None
                font_size = None
                bold = False
                italic = False
            
            # Clear all runs
            for run in paragraph.runs:
                run.text = ""
            
            # Add new run with replaced text
            if paragraph.runs:
                paragraph.runs[0].text = full_text
            else:
                new_run = paragraph.add_run(full_text)
                if font_name:
                    new_run.font.name = font_name
                if font_size:
                    new_run.font.size = font_size
                new_run.font.bold = bold
                new_run.font.italic = italic
    
    def _format_value(self, value: Any, field_name: str) -> str:
        """Format a value based on its type."""
        if value is None:
            return ""
        
        name_lower = field_name.lower()
        
        # Format dates
        if isinstance(value, datetime):
            return value.strftime("%d.%m.%Y")
        
        # Format numbers with thousands separator
        if isinstance(value, (int, float)):
            if any(x in name_lower for x in ["amount", "summa", "sum", "qiymat"]):
                if isinstance(value, float):
                    return f"{value:,.2f}".replace(",", " ")
                return f"{value:,}".replace(",", " ")
            return str(value)
        
        return str(value)
    
    def fill_xlsx_template(
        self,
        template_path: Path,
        field_values: dict[str, Any],
        output_path: Path | None = None,
    ) -> Path:
        """Fill an Excel template with provided field values."""
        wb = load_workbook(template_path)
        ws = wb.active
        
        # Iterate through all cells
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    for match in self.FIELD_PATTERN.finditer(cell.value):
                        field_name = match.group(1)
                        placeholder = match.group(0)
                        
                        if field_name in field_values:
                            value = self._format_value(field_values[field_name], field_name)
                            cell.value = cell.value.replace(placeholder, str(value))
        
        # Generate output path if not provided
        if output_path is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_name = f"{template_path.stem}_{timestamp}.xlsx"
            output_path = self.output_dir / output_name
        
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        
        return output_path
    
    def calculate_match_score(
        self,
        template_fields: list[str],
        available_data: dict[str, Any],
    ) -> dict:
        """
        Calculate how well available data matches a template.
        
        Returns:
            dict with match_score (0-1), matched_fields, missing_fields
        """
        required_fields = template_fields
        available_keys = set(available_data.keys())
        
        # Normalize field names for comparison
        normalized_available = {self._normalize_field(k): k for k in available_keys}
        
        matched = []
        missing = []
        
        for field in required_fields:
            normalized = self._normalize_field(field)
            if normalized in normalized_available:
                matched.append(field)
            else:
                # Try fuzzy matching
                for norm_avail, orig in normalized_available.items():
                    if self._fuzzy_match(normalized, norm_avail):
                        matched.append(field)
                        break
                else:
                    missing.append(field)
        
        total = len(required_fields)
        score = len(matched) / total if total > 0 else 0
        
        return {
            "match_score": round(score, 2),
            "matched_fields": matched,
            "missing_fields": missing,
            "total_fields": total,
        }
    
    def _normalize_field(self, field: str) -> str:
        """Normalize field name for comparison."""
        return field.lower().replace("_", "").replace("-", "").replace(" ", "")
    
    def _fuzzy_match(self, field1: str, field2: str) -> bool:
        """Check if two field names are similar enough."""
        # Simple substring matching
        if field1 in field2 or field2 in field1:
            return True
        
        # Common translations
        translations = {
            "company": ["korxona", "firma", "tashkilot"],
            "name": ["nomi", "ism"],
            "date": ["sana", "kun"],
            "amount": ["summa", "miqdor"],
            "address": ["manzil"],
            "phone": ["telefon"],
        }
        
        for eng, uzb_list in translations.items():
            if eng in field1:
                return any(u in field2 for u in uzb_list)
            for uzb in uzb_list:
                if uzb in field1:
                    return eng in field2
        
        return False


# Global instance
template_service = TemplateService()
