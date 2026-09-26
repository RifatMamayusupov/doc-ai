from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import pandas as pd
import os

def create_presentation():
    try:
        # Initialize Presentation
        prs = Presentation()
        
        # --- Slide 1: Title Slide ---
        slide_layout = prs.slide_layouts[0] # Title Slide
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "Loyiha Tahlili: Qorabog' Suv Ombori"
        subtitle.text = "Avtomatik Generatsiya qilingan Hisobot\nExcel Agent tomonidan tayyorlandi"
        
        print("Created Title Slide")

        # --- Slide 2: Overview ---
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = "Loyiha Haqida Qisqacha"
        
        # Read data stats
        df = pd.read_excel('cleaned_data.xlsx')
        total_rows = len(df)
        total_cols = len(df.columns)
        
        text_frame = content.text_frame
        p = text_frame.add_paragraph()
        p.text = f"Fayl nomi: complex_structure_test.xlsx"
        p.level = 0
        
        p = text_frame.add_paragraph()
        p.text = f"Tozalangan ma'lumotlar hajmi: {total_rows} qator"
        p.level = 0
        
        p = text_frame.add_paragraph()
        p.text = "Asosiy ish turlari:"
        p.level = 0
        
        p = text_frame.add_paragraph()
        p.text = "• Tuproq ishlari"
        p.level = 1
        p = text_frame.add_paragraph()
        p.text = "• Beton ishlari"
        p.level = 1
        p = text_frame.add_paragraph()
        p.text = "• Mexanizatsiya xarajatlari"
        p.level = 1

        print("Created Overview Slide")

        # --- Slide 3: Visualization ---
        if os.path.exists('pie_chart.png'):
            slide_layout = prs.slide_layouts[5] # Title Only
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            title.text = "O'lchov Birliklari Taqsimoti"
            
            # Add image
            left = Inches(1)
            top = Inches(1.5)
            height = Inches(5.5)
            slide.shapes.add_picture('pie_chart.png', left, top, height=height)
            print("Created Visualization Slide")
        
        # --- Slide 4: Data Table (Top 5 rows) ---
        slide_layout = prs.slide_layouts[5] # Title Only
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "Ma'lumotlardan Namuna (Top 5)"
        
        # Prepare data for table
        top_df = df.head(5)
        # Select only first 4 columns to fit in slide
        top_df = top_df.iloc[:, :4] 
        
        rows, cols = top_df.shape
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(0.8)
        
        table = slide.shapes.add_table(rows+1, cols, left, top, width, height).table
        
        # Set column headers
        for i, col_name in enumerate(top_df.columns):
            cell = table.cell(0, i)
            cell.text = str(col_name)
            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(10)
            
        # Fill rows
        for i in range(rows):
            for j in range(cols):
                cell = table.cell(i+1, j)
                val = top_df.iloc[i, j]
                # Truncate long text
                text_val = str(val)
                if len(text_val) > 50:
                    text_val = text_val[:47] + "..."
                cell.text = text_val
                cell.text_frame.paragraphs[0].font.size = Pt(9)

        print("Created Table Slide")

        # Save presentation
        output_file = 'Qorabog_Hisobot.pptx'
        prs.save(output_file)
        print(f"Successfully saved presentation to {output_file}")

    except Exception as e:
        print(f"Error creating presentation: {e}")

if __name__ == "__main__":
    create_presentation()
